"""Выгрузка базы в Excel."""
from __future__ import annotations

import io
from datetime import datetime

from aiogram import F, Router
from aiogram.types import BufferedInputFile, Message
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from ..access import Private, Role
from ..dates import next_occurrence
from ..db import DB
from ..notifier import Notifier
from ..ui import BTN_EXPORT

router = Router(name="export")
router.message.filter(Private(), Role("admin", "hr"))

COLS = {
    "employee": [("ФИО", "full_name"), ("Дата рождения", "_bd"), ("Ближайший ДР", "_next"),
                 ("Должность", "position"), ("Отдел", "department"), ("Формат", "gift_format"),
                 ("Характеристика", "characteristic"), ("Публично", "_pub"), ("Телефон", "phone"),
                 ("Комментарий", "comment"), ("Активен", "_act")],
    "partner": [("ФИО", "full_name"), ("Дата рождения", "_bd"), ("Ближайший ДР", "_next"),
                ("Компания", "company"), ("Категория", "category"), ("Ответственный", "responsible_name"),
                ("Дарим", "gift_format"), ("Как", "congrat_method"), ("Кто поздравляет", "congratulator"),
                ("Телефон", "phone"), ("Мессенджер", "messenger"), ("Комментарий", "comment"),
                ("Активен", "_act")],
}


def build_xlsx(rows: list[dict], today, feb29: str) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for kind, title in (("employee", "Сотрудники"), ("partner", "Партнёры")):
        ws = wb.create_sheet(title)
        cols = COLS[kind]
        ws.append([c[0] for c in cols])
        for c in ws[1]:
            c.font = Font(bold=True)
        for p in sorted((r for r in rows if r["kind"] == kind), key=lambda r: (r["birth_month"], r["birth_day"])):
            out = []
            for _, key in cols:
                if key == "_bd":
                    y = p["birth_year"] or "????"
                    out.append(f"{p['birth_day']:02d}.{p['birth_month']:02d}.{y}")
                elif key == "_next":
                    out.append(next_occurrence(p["birth_day"], p["birth_month"], today, feb29))
                elif key == "_pub":
                    out.append("да" if p["public_congrats"] else "нет")
                elif key == "_act":
                    out.append("да" if p["is_active"] else "архив")
                else:
                    out.append(p.get(key))
            ws.append(out)
        for i, (name, key) in enumerate(cols, 1):
            ws.column_dimensions[get_column_letter(i)].width = 40 if key in ("characteristic", "comment") else 20
            if key == "_next":
                for cell in ws[get_column_letter(i)][1:]:
                    cell.number_format = "DD.MM.YYYY"
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@router.message(F.text == BTN_EXPORT)
async def export(message: Message, db: DB, notifier: Notifier):
    rows = await db.persons(active=None)
    data = build_xlsx(rows, notifier.today(), await notifier.feb29())
    name = f"birthdays_{datetime.now():%Y-%m-%d}.xlsx"
    await message.answer_document(BufferedInputFile(data, filename=name),
                                  caption=f"Выгрузка: {len(rows)} записей")
